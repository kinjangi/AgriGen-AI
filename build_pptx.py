import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_slide_layout = prs.slide_layouts[6]

    # Colors
    BG_DARK = RGBColor(15, 23, 42)      # #0f172a
    CARD_BG = RGBColor(30, 41, 59)      # #1e293b
    CARD_BORDER = RGBColor(51, 65, 85)  # #334155
    ACCENT_GREEN = RGBColor(52, 211, 153) # #34d399
    ACCENT_BLUE = RGBColor(56, 189, 248)  # #38bdf8
    ACCENT_YELLOW = RGBColor(250, 204, 21) # #facc15
    TEXT_WHITE = RGBColor(248, 250, 252) # #f8fafc
    TEXT_MUTED = RGBColor(148, 163, 184) # #94a3b8

    def set_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = BG_DARK

    def add_header(slide, title_text, presenter_text, category="PHASE 1 PROJECT REVIEW"):
        # Header category
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_c = cat_box.text_frame
        tf_c.word_wrap = True
        p_c = tf_c.paragraphs[0]
        p_c.text = category.upper()
        p_c.font.size = Pt(10)
        p_c.font.bold = True
        p_c.font.color.rgb = ACCENT_BLUE

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(8.5), Inches(0.8))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE

        # Presenter Badge
        pres_box = slide.shapes.add_textbox(Inches(9.2), Inches(0.65), Inches(3.3), Inches(0.8))
        tf_p = pres_box.text_frame
        tf_p.word_wrap = True
        p_p = tf_p.paragraphs[0]
        p_p.alignment = PP_ALIGN.RIGHT
        p_p.text = f"🎤 Presenter:\n{presenter_text}"
        p_p.font.size = Pt(11)
        p_p.font.bold = True
        p_p.font.color.rgb = ACCENT_GREEN

    def add_card(slide, left, top, width, height, title, items, border_color=CARD_BORDER):
        # Card Background
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)

        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.2)
        tf.margin_bottom = Inches(0.2)

        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.size = Pt(15)
        p0.font.bold = True
        p0.font.color.rgb = ACCENT_BLUE
        p0.space_after = Pt(10)

        for item in items:
            p = tf.add_paragraph()
            p.text = f"•  {item}"
            p.font.size = Pt(12)
            p.font.color.rgb = TEXT_WHITE
            p.space_after = Pt(6)

    # ------------------- SLIDE 1: Title Slide -------------------
    slide1 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide1)
    
    # Large Title Card
    shape = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.0), Inches(11.333), Inches(5.5))
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD_BG
    shape.line.color.rgb = ACCENT_GREEN
    shape.line.width = Pt(2)
    tf1 = shape.text_frame
    tf1.word_wrap = True
    tf1.margin_left = Inches(0.5)
    tf1.margin_top = Inches(0.5)

    p = tf1.paragraphs[0]
    p.text = "AgriGen AI: Intelligent Multi-Agent Farming Ecosystem"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    p.space_after = Pt(8)

    p2 = tf1.add_paragraph()
    p2.text = "Phase 1 Comprehensive Project Review Presentation"
    p2.font.size = Pt(18)
    p2.font.color.rgb = TEXT_WHITE
    p2.space_after = Pt(24)

    p3 = tf1.add_paragraph()
    p3.text = "👥 Team Members & Roles (6-Member Team):"
    p3.font.size = Pt(14)
    p3.font.bold = True
    p3.font.color.rgb = ACCENT_BLUE
    p3.space_after = Pt(10)

    team_members = [
        "1. Member 1 (Team Lead & Product Architect) — System Architecture, Phase 1 Progress & Strategy",
        "2. Member 2 (Frontend & UI/UX Lead) — React/Vite Dashboard, Multilingual Voice & Scanner UI",
        "3. Member 3 (Backend & Orchestration Lead) — FastAPI Microservices, Async SSE Streaming & Routes",
        "4. Member 4 (Multi-Agent & RAG Specialist) — LangGraph Supervisor, FAISS Vector Knowledge Base",
        "5. Member 5 (Machine Learning & Vision Lead) — PyTorch LSTM (Weather & Market) & Disease Vision",
        "6. Member 6 (DevOps & Database Integration Lead) — PostgreSQL/SQLite Data Models, Docker & Scripts"
    ]
    for tm in team_members:
        p_m = tf1.add_paragraph()
        p_m.text = tm
        p_m.font.size = Pt(11)
        p_m.font.color.rgb = TEXT_WHITE
        p_m.space_after = Pt(4)

    slide1.notes_slide.notes_text_frame.text = (
        "SPEAKER SCRIPT (Member 1 - Team Lead):\n"
        "Good morning respected mentor and review panel. Today, our 6-member team is proud to present the Phase 1 Project Review for 'AgriGen AI: Intelligent Multi-Agent Farming Ecosystem'. "
        "AgriGen AI addresses critical agricultural challenges through deep learning, multi-agent LLM orchestration, computer vision, and real-time market/weather forecasting. "
        "During this presentation, each of our 6 team members will detail their specific engineering contributions."
    )

    # ------------------- SLIDE 2: Problem & Objectives -------------------
    slide2 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide2)
    add_header(slide2, "Problem Statement & Project Objectives", "Member 1 (Team Leader)")

    add_card(slide2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "Agricultural Challenges", [
        "Price Volatility: Farmers sell crops blindly without historical trend forecasts.",
        "Micro-Climate Surprises: Generic weather forecasts lack farm-level precision.",
        "Crop Disease Losses: Delayed leaf pest/disease diagnosis leads to total crop destruction.",
        "Information Silos & Language Barrier: Research papers are inaccessible; rural farmers struggle with text-only English apps."
    ], ACCENT_YELLOW)

    add_card(slide2, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), "AgriGen AI Objectives", [
        "Multi-Agent Intelligence: Deploy specialist agents using LangGraph & Llama-3.",
        "LSTM Predictive Engines: Provide 7-day weather & crop price forecasts with peak profit alerts.",
        "Instant Vision Diagnostics: Computer vision to scan leaves and prescribe actionable remedies.",
        "RAG Agricultural Advisory: FAISS vector database indexed with verified research papers.",
        "Multilingual Voice Interface: Native English, Hindi, and Telugu speech interaction."
    ], ACCENT_GREEN)

    slide2.notes_slide.notes_text_frame.text = (
        "SPEAKER SCRIPT (Member 1):\n"
        "Modern farmers face compounding risks: volatile market prices, localized crop diseases, and unpredicted weather shifts. "
        "Our objective with AgriGen AI is to build a unified multi-agent platform that gives every farmer a hyper-personalized digital agronomist right in their pocket."
    )

    # ------------------- SLIDE 3: System Architecture -------------------
    slide3 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide3)
    add_header(slide3, "System Architecture & High-Level Workflow", "Member 1 (Team Leader)")

    add_card(slide3, Inches(0.8), Inches(1.6), Inches(4.5), Inches(5.2), "Core Architectural Pillars", [
        "Decoupled Microservices: React Frontend + FastAPI Async Backend.",
        "Stateful Multi-Agent Graph: LangGraph Supervisor manages node state transitions.",
        "Hybrid Predictive Analytics: PyTorch LSTMs + Computer Vision + FAISS Vector DB.",
        "Dual Storage Strategy: PostgreSQL for relational persistence, Redis for caching & state memory."
    ], ACCENT_BLUE)

    if os.path.exists("system_architecture.png"):
        slide3.shapes.add_picture("system_architecture.png", Inches(5.5), Inches(1.6), Inches(7.0), Inches(5.2))

    slide3.notes_slide.notes_text_frame.text = (
        "SPEAKER SCRIPT (Member 1):\n"
        "Here is our system architecture diagram. User queries enter via our React frontend or voice widget, reach FastAPI, and are dispatched to the LangGraph Supervisor. "
        "The supervisor dynamically delegates execution to specialized agent nodes. I will now hand over to Member 2 to present the Frontend Dashboard."
    )

    # ------------------- SLIDE 4: Frontend Development -------------------
    slide4 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide4)
    add_header(slide4, "Frontend Dashboard & UI/UX Architecture", "Member 2 (Frontend Specialist)")

    add_card(slide4, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "User Interface & Experience", [
        "Vite + React Core: Built for ultra-fast startup and smooth client-side interactions.",
        "Custom Agricultural Design System: Earthy dark palette, glassmorphism, tailored typography.",
        "Interactive Analytics: Dynamic Chart.js integration for price trends & moisture metrics.",
        "Location Picker Modal: Interactive GPS coordinate selector and field selector bar."
    ], ACCENT_GREEN)

    add_card(slide4, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), "Multilingual Assistant Widget", [
        "Floating Chat Widget: Persistent bottom-right drawer supporting rich Markdown responses.",
        "Voice Input / Output: Browser STT/TTS engine supporting English, Hindi, and Telugu.",
        "Drag-and-Drop Image Uploader: Instant leaf photo scanning for disease identification.",
        "Persistent User Memory: LocalStorage session preservation across chat resets."
    ], ACCENT_BLUE)

    slide4.notes_slide.notes_text_frame.text = (
        "SPEAKER SCRIPT (Member 2):\n"
        "Respected panel, as the Frontend Lead, my focus was building a responsive, accessible interface for farmers. "
        "We developed a React/Vite dashboard featuring custom charts, an interactive map location picker, and a floating AI assistant widget that supports Hindi and Telugu voice input."
    )

    # ------------------- SLIDE 5: Backend & API -------------------
    slide5 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide5)
    add_header(slide5, "Backend Microservices & Orchestration Layer", "Member 3 (Backend Engineer)")

    add_card(slide5, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "FastAPI Microservices Architecture", [
        "Async Python Engine: Built on FastAPI and Uvicorn for high-concurrency request handling.",
        "Modular V1 API Routing: Clean separation across auth, weather, market, disease, and report routes.",
        "Pydantic Validation: Strict schema validation for all API inputs and structured outputs.",
        "CORS & Security Controls: Configured middleware with environment variable isolation."
    ], ACCENT_BLUE)

    add_card(slide5, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), "Real-Time Streaming & Orchestrator", [
        "SSE Streaming Endpoint (`/api/v1/orchestrator/chat`): Real-time token streaming to frontend.",
        "Stateful Session Management: Tracks conversation histories and user preferences in memory.",
        "WeasyPrint PDF Generator: Asynchronous trigger for weekly summary PDF reports.",
        "Error Middleware: Graceful exception handling and fallback logging."
    ], ACCENT_GREEN)

    slide5.notes_slide.notes_text_frame.text = (
        "SPEAKER SCRIPT (Member 3):\n"
        "As the Backend Engineer, I designed our asynchronous FastAPI core. "
        "Our streaming orchestrator endpoint utilizes Server-Sent Events (SSE), allowing AI agent responses to stream incrementally to the farmer without blocking the server."
    )

    # ------------------- SLIDE 6: LangGraph & RAG -------------------
    slide6 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide6)
    add_header(slide6, "LangGraph Multi-Agent Router & RAG Advisory Engine", "Member 4 (Multi-Agent Lead)")

    add_card(slide6, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "LangGraph Supervisor Workflow", [
        "Directed Acyclic Graph (DAG): Implements state-machine routing powered by LLMs.",
        "Intent Router: Evaluates incoming prompt and context to assign the optimal agent.",
        "Specialist Nodes: Dedicated nodes for Weather, Market, Disease, and Knowledge RAG.",
        "State Preservation: State dictionary maintains thread history and agent decisions."
    ], ACCENT_GREEN)

    add_card(slide6, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), "FAISS RAG Knowledge Engine", [
        "Curated Agriculture Dataset: Indexed ICAR research papers, soil health data & subsidies.",
        "Vector Embeddings: Uses HuggingFace all-MiniLM-L6-v2 embeddings for fast retrieval.",
        "Similarity Thresholding: Mitigates LLM hallucinations by retrieving top-k verified context chunks.",
        "Grounded Advisory: Formulates actionable farmer responses strictly supported by retrieved facts."
    ], ACCENT_YELLOW)

    slide6.notes_slide.notes_text_frame.text = (
        "SPEAKER SCRIPT (Member 4):\n"
        "My core contribution was building the multi-agent decision engine using LangGraph and RAG. "
        "The Supervisor evaluates intent—for example, if a query mentions 'fertilizer advice', it routes to our FAISS vector store containing official ICAR research."
    )

    # ------------------- SLIDE 7: Predictive Analytics & ML -------------------
    slide7 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide7)
    add_header(slide7, "Predictive Analytics & Deep Learning Models", "Member 5 (Machine Learning Lead)")

    add_card(slide7, Inches(0.8), Inches(1.6), Inches(3.7), Inches(5.2), "Weather LSTM Model", [
        "PyTorch Architecture: Sequential LSTM network.",
        "Inputs: Historical temp, humidity, pressure, rainfall.",
        "Output: 7-day micro-climate precipitation & humidity trend.",
        "Accuracy: High R² score on local weather datasets."
    ], ACCENT_BLUE)

    add_card(slide7, Inches(4.8), Inches(1.6), Inches(3.7), Inches(5.2), "Market LSTM Model", [
        "Price Forecasting: Predicts mandi prices for major crops.",
        "Profitability Peak Alert: Identifies optimal selling window.",
        "Mandi Integration: Dynamic price tracking for regional markets."
    ], ACCENT_GREEN)

    add_card(slide7, Inches(8.8), Inches(1.6), Inches(3.7), Inches(5.2), "Leaf Disease Vision", [
        "CNN / Vision Classifier: Scans crop leaves for early pest/disease detection.",
        "Confidence Scoring: Outputs disease probability & severity level.",
        "Treatment Link: Directly maps detected disease to organic & chemical treatments."
    ], ACCENT_YELLOW)

    slide7.notes_slide.notes_text_frame.text = (
        "SPEAKER SCRIPT (Member 5):\n"
        "As the ML Lead, I developed our PyTorch models. "
        "We have two trained LSTMs—one for 7-day weather forecasting and one for crop price prediction—alongside our Computer Vision module that identifies crop leaf diseases from images."
    )

    # ------------------- SLIDE 8: Database & DevOps -------------------
    slide8 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide8)
    add_header(slide8, "Database Architecture, DevOps & Infrastructure", "Member 6 (DevOps & Database Lead)")

    add_card(slide8, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "Data Storage & Schema Design", [
        "SQLAlchemy ORM: Unified data access layer across environments.",
        "Production PostgreSQL & Local SQLite: Seamless fallback mechanism (`agrigen.db`).",
        "Key Schemas: User profiles, farm plots, historical weather, market rates, diagnostic logs.",
        "Redis Caching: Fast session persistence and rate-limiting."
    ], ACCENT_GREEN)

    add_card(slide8, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), "Containerization & Local Dev", [
        "Docker & Docker Compose: Containerized backend, frontend, and database services.",
        "One-Click PowerShell Launcher (`run_local.ps1`): Launches backend and frontend concurrently.",
        "Virtual Environment Management: Custom virtual environment with patched hashlib dependencies.",
        "Production NGINX Config: Reverse proxy setup with SSL readiness."
    ], ACCENT_BLUE)

    slide8.notes_slide.notes_text_frame.text = (
        "SPEAKER SCRIPT (Member 6):\n"
        "My role focused on database design, devops, and local launcher orchestration. "
        "I implemented an ORM layer supporting PostgreSQL and local SQLite, containerized the platform using Docker, and automated local developer launch via `run_local.ps1`."
    )

    # ------------------- SLIDE 9: Phase 1 Progress -------------------
    slide9 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide9)
    add_header(slide9, "Phase 1 Implementation Progress & Key Milestones", "Member 1 (Team Leader)")

    add_card(slide9, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "Completed Deliverables (Phase 1)", [
        "✅ Functional Multi-Agent Workflow: LangGraph Supervisor with specialist agent nodes.",
        "✅ Responsive React Dashboard: Operational charts, location picker, and floating chat widget.",
        "✅ FastAPI Backend: Async REST routes, SSE streaming endpoint, and WeasyPrint PDF reports.",
        "✅ Trained Machine Learning Baseline Models: PyTorch LSTMs & Leaf Disease classifier.",
        "✅ FAISS RAG Pipeline: Ingested agricultural research papers for advisory."
    ], ACCENT_GREEN)

    add_card(slide9, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), "Testing & Verification Progress", [
        "End-to-End Local Execution: Successfully tested via `run_local.ps1`.",
        "Multilingual Voice Test: Verified STT/TTS in English, Hindi, and Telugu.",
        "API Test Coverage: Verified all `/api/v1/` routes with Swagger UI (`/docs`).",
        "Robustness Patching: Resolved Windows native C-extension DLL fallback issues."
    ], ACCENT_BLUE)

    slide9.notes_slide.notes_text_frame.text = (
        "SPEAKER SCRIPT (Member 1):\n"
        "We are pleased to report that Phase 1 objectives have been successfully achieved! "
        "Our frontend dashboard, FastAPI backend, multi-agent supervisor, PyTorch models, and RAG knowledge base are fully integrated and running smoothly."
    )

    # ------------------- SLIDE 10: Technical Challenges -------------------
    slide10 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide10)
    add_header(slide10, "Key Technical Challenges & Engineering Solutions", "Member 3 & Member 4")

    add_card(slide10, Inches(0.8), Inches(1.6), Inches(3.7), Inches(5.2), "LLM Hallucinations", [
        "Challenge: Generic LLMs gave unsafe agricultural advice.",
        "Solution: Implemented FAISS RAG vector retriever with strict confidence scoring."
    ], ACCENT_YELLOW)

    add_card(slide10, Inches(4.8), Inches(1.6), Inches(3.7), Inches(5.2), "Multi-Agent Latency", [
        "Challenge: Sequential agent calls caused response lag.",
        "Solution: Enabled Server-Sent Events (SSE) streaming for real-time response rendering."
    ], ACCENT_BLUE)

    add_card(slide10, Inches(8.8), Inches(1.6), Inches(3.7), Inches(5.2), "Native C DLL Conflict", [
        "Challenge: Windows App Control blocked `xxhash` DLL in `langgraph`.",
        "Solution: Patched `langgraph` with pure-python `hashlib` fallback."
    ], ACCENT_GREEN)

    slide10.notes_slide.notes_text_frame.text = (
        "SPEAKER SCRIPT (Member 3 & 4):\n"
        "During Phase 1, we tackled critical engineering hurdles. "
        "We solved LLM hallucinations using FAISS RAG grounding, eliminated latency using SSE streaming, and bypassed Windows DLL restrictions with a custom hashlib patch."
    )

    # ------------------- SLIDE 11: Future Scope -------------------
    slide11 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide11)
    add_header(slide11, "Phase 2 Roadmap & Future Scope", "Member 5 (Machine Learning Lead)")

    add_card(slide11, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "Phase 2 Feature Roadmap", [
        "Edge Model Optimization: Quantize disease CNN to ONNX for offline mobile inference.",
        "IoT Smart Valve Integration: Direct relay control for autonomous irrigation based on weather LSTMs.",
        "Drone Orthomosaic Scanning: Processing aerial multispectral imagery for wide-area field analysis.",
        "Mandi Direct Trading: Automated market order execution at peak profitability windows."
    ], ACCENT_BLUE)

    add_card(slide11, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), "Scalability & Deployment", [
        "Kubernetes Auto-Scaling: Elastic cluster scaling for CPU-intensive PyTorch inference.",
        "Satellite Telemetry Integration: Real-time Sentinel-2 NDVI crop health index ingestion.",
        "Vernacular Voice Expansion: Adding regional Indian dialects (Tamil, Kannada, Marathi)."
    ], ACCENT_GREEN)

    slide11.notes_slide.notes_text_frame.text = (
        "SPEAKER SCRIPT (Member 5):\n"
        "Looking ahead to Phase 2, our roadmap includes deploying edge vision models for offline mobile use, integrating IoT smart irrigation controllers, and incorporating satellite crop telemetry."
    )

    # ------------------- SLIDE 12: Work Breakdown Matrix -------------------
    slide12 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide12)
    add_header(slide12, "Team Work Breakdown & Individual Contribution Matrix", "Member 6 (DevOps Lead)")

    add_card(slide12, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2), "6-Member Deliverables Breakdown", [
        "Member 1 (Team Leader): Overall architecture design, supervisor workflow configuration, Phase 1 roadmap & review lead.",
        "Member 2 (Frontend Specialist): React/Vite UI components, location picker, dynamic Chart.js dashboards & voice widget.",
        "Member 3 (Backend Engineer): FastAPI orchestrator, SSE streaming endpoints, API routing structure & WeasyPrint PDF reports.",
        "Member 4 (Multi-Agent & RAG Specialist): LangGraph agent state machine, FAISS vector embeddings, agricultural RAG pipeline.",
        "Member 5 (Machine Learning Lead): PyTorch LSTM training (weather & crop price forecasting) & leaf disease vision classifier.",
        "Member 6 (DevOps & Database Lead): PostgreSQL/SQLite data schemas, Docker Compose orchestration, Redis cache, local run script."
    ], ACCENT_GREEN)

    slide12.notes_slide.notes_text_frame.text = (
        "SPEAKER SCRIPT (Member 6):\n"
        "This matrix details the clear distribution of responsibilities across our 6 team members. "
        "Every member has taken ownership of a major component of the AgriGen AI stack."
    )

    # ------------------- SLIDE 13: Live Demo Showcase -------------------
    slide13 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide13)
    add_header(slide13, "Live Demonstration & Workflow Execution", "Member 2 & Team")

    if os.path.exists("agent_workflow.png"):
        slide13.shapes.add_picture("agent_workflow.png", Inches(0.8), Inches(1.6), Inches(6.5), Inches(5.2))

    add_card(slide13, Inches(7.5), Inches(1.6), Inches(5.0), Inches(5.2), "Demo Highlights", [
        "1. Dynamic Dashboard: Location picker updates local mandi prices & weather forecasts.",
        "2. Voice Input: Querying crop advisory in Telugu/Hindi via voice button.",
        "3. Leaf Scanner: Uploading leaf photo -> Instant disease detection & organic remedies.",
        "4. PDF Generation: One-click export of farmer summary report."
    ], ACCENT_BLUE)

    slide13.notes_slide.notes_text_frame.text = (
        "SPEAKER SCRIPT (Member 2):\n"
        "We demonstrate our live execution flow: from farmer input to LangGraph routing, PyTorch inference, FAISS retrieval, and instant UI rendering."
    )

    # ------------------- SLIDE 14: Conclusion & Q&A -------------------
    slide14 = prs.slides.add_slide(blank_slide_layout)
    set_background(slide14)
    add_header(slide14, "Conclusion & Panel Q&A Defense Strategy", "Member 1 & Team")

    add_card(slide14, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "Summary of Phase 1 Readiness", [
        "Fully Operational Prototype: All core modules integrated and tested locally.",
        "High Technical Rigor: Deep learning, microservices, and RAG working in unison.",
        "Social Impact: Empowers rural farmers with localized, accessible AI advisory.",
        "Prepared for Phase 2: Solid foundation established for IoT & Edge expansion."
    ], ACCENT_GREEN)

    add_card(slide14, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), "Thank You! Questions & Discussion", [
        "We welcome questions and feedback from our respected mentor and review panel.",
        "Project Repository: agriGen_AI",
        "Team Members: Members 1, 2, 3, 4, 5, 6",
        "Status: Ready for Phase 1 Sign-Off!"
    ], ACCENT_YELLOW)

    slide14.notes_slide.notes_text_frame.text = (
        "SPEAKER SCRIPT (Member 1 & All):\n"
        "In conclusion, AgriGen AI represents a major step forward in accessible agricultural intelligence. "
        "Thank you for your time and guidance. We are now open for questions and feedback from the panel."
    )

    output_pptx = os.path.join(os.getcwd(), "AgriGen_AI_Phase1_Project_Review.pptx")
    prs.save(output_pptx)
    print(f"PowerPoint Presentation successfully created at: {output_pptx}")

if __name__ == "__main__":
    create_presentation()
